"""Pulling a hierarchy of sections out of a .docx policy document.

Word heading styles are the hierarchy, and in a policy document they are
reliable in a way almost nothing else about the file is: the numbering scheme
varies, the fonts vary, the tables vary, but somebody applied Heading 1 and
Heading 2 because that is how the table of contents gets built.

PDF is deliberately not handled. Policy in a SharePoint library is nearly
always the Word original, the PDF is a rendering of it, and PDF extraction is
a different problem with a much worse failure mode -- silently reordered
columns rather than a clean error.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

# Leading section number in a heading: "5.1 SAR Filing Deadlines" ->
# ("5.1", "SAR Filing Deadlines"). Policies that do not number their headings
# get a positional number instead, so the citation still points somewhere.
_NUMBERED_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+)*\.?)\s+(.*\S)\s*$")
_HEADING_STYLE_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


class ExtractionError(RuntimeError):
    """The document could not be read as a policy."""


@dataclass
class ExtractedSection:
    number: str
    heading: str
    level: int
    heading_path: tuple[str, ...]
    paragraphs: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(self.paragraphs).strip()


@dataclass
class ExtractedDocument:
    title: str
    source_path: str
    sections: list[ExtractedSection]


def _heading_level(style_name: str | None) -> int | None:
    if not style_name:
        return None
    match = _HEADING_STYLE_RE.match(style_name.strip())
    if match:
        return int(match.group(1))
    if style_name.strip().lower() == "title":
        return 0
    return None


def _split_heading(raw: str, fallback_number: str) -> tuple[str, str]:
    match = _NUMBERED_HEADING_RE.match(raw)
    if match:
        return match.group(1).rstrip("."), match.group(2)
    return fallback_number, raw.strip()


def extract_docx(path: str | Path) -> ExtractedDocument:
    """Read a .docx into a flat list of sections, each knowing its ancestry."""
    try:
        import docx  # noqa: PLC0415 - optional dependency, imported on use
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise ExtractionError(
            "Reading .docx needs python-docx. Install it with "
            "`uv pip install python-docx`."
        ) from exc

    source = Path(path)
    if not source.exists():
        raise ExtractionError(f"No such file: {source}")

    try:
        document = docx.Document(str(source))
    except Exception as exc:
        raise ExtractionError(f"{source.name} could not be opened as .docx: {exc}") from exc

    title = source.stem
    sections: list[ExtractedSection] = []
    # Stack of (level, "number heading") for the ancestry of the current section.
    ancestry: list[tuple[int, str]] = []
    preamble: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        level = _heading_level(getattr(paragraph.style, "name", None))
        if level is None:
            if sections:
                sections[-1].paragraphs.append(text)
            else:
                # Text before the first heading: front matter, kept so that a
                # one-page policy with no headings is not silently empty.
                preamble.append(text)
            continue

        if level == 0:
            title = text
            continue

        number, heading = _split_heading(text, fallback_number=str(len(sections) + 1))
        while ancestry and ancestry[-1][0] >= level:
            ancestry.pop()
        path_parts = tuple(part for _, part in ancestry)
        ancestry.append((level, f"{number} {heading}".strip()))

        sections.append(
            ExtractedSection(
                number=number,
                heading=heading,
                level=level,
                heading_path=(title,) + path_parts + (f"{number} {heading}".strip(),),
            )
        )

    if preamble and not sections:
        sections.append(
            ExtractedSection(
                number="1",
                heading=title,
                level=1,
                heading_path=(title,),
                paragraphs=preamble,
            )
        )

    if not sections:
        raise ExtractionError(f"{source.name} produced no text; is it empty?")

    return ExtractedDocument(title=title, source_path=str(source), sections=sections)
